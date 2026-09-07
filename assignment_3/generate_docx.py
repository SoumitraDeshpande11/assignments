"""Convert DOCUMENTATION.md into DOCUMENTATION.docx (minimal styling)."""
import re

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DOC_PATH = "DOCUMENTATION.md"
OUT_PATH = "DOCUMENTATION.docx"
IMG_PATH = "assets/screenshot.png"

ACCENT = RGBColor(0x2D, 0x6A, 0x4F)


def add_runs(paragraph, text):
    for part in re.split(r"(\*\*.*?\*\*|`[^`]+`)", text):
        if not part:
            continue
        run = paragraph.add_run()
        if part.startswith("**"):
            run.text = part[2:-2]
            run.bold = True
        elif part.startswith("`"):
            run.text = part[1:-1]
            run.font.name = "Menlo"
            run.font.size = Pt(9.5)
        else:
            run.text = part


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            target = table.cell(r, c)
            target.paragraphs[0].text = ""
            add_runs(target.paragraphs[0], cell)
            if r == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E9F2EE")
                target._tc.get_or_add_tcPr().append(shading)
                for run in target.paragraphs[0].runs:
                    run.bold = True


def is_separator(row):
    return all(set(c) <= set("-: ") for c in row)


def build():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    with open(DOC_PATH) as f:
        lines = f.read().splitlines()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line or line == "---":
            i += 1
            continue
        if line.startswith("!["):
            doc.add_picture(IMG_PATH, width=Inches(2.6))
            i += 1
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            heading = doc.add_heading("", level=min(level, 3))
            add_runs(heading, line.lstrip("# "))
            for run in heading.runs:
                run.font.color.rgb = ACCENT
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].split("|")[1:-1]]
                if not is_separator(cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue
        if line.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(block))
            run.font.name = "Menlo"
            run.font.size = Pt(9.5)
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line[2:])
            i += 1
            continue
        p = doc.add_paragraph()
        add_runs(p, line)
        i += 1

    doc.save(OUT_PATH)
    print(f"Saved {OUT_PATH}")


if __name__ == "__main__":
    build()
